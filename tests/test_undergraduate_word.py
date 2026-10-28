from __future__ import annotations
import copy
import json
from pathlib import Path

import pytest
from docx import Document
from docx.shared import Inches
from PIL import Image
import pypdfium2 as pdfium

from unified_paper_orchestrator.undergraduate_word import (
    SCHEMA, brief_digest, digest, inspect_docx, is_word_workflow, paper_workflow,
    validate_word_delivery, write_context, _layout, _convert_tex,
)


@pytest.fixture
def delivery(tmp_path):
    # Mechanical receipt fixture, never a claim that a real thesis was reviewed.
    root = tmp_path
    paper = root / "papers" / "ug"
    submit = paper / "submit"
    submit.mkdir(parents=True)
    out = paper / "delivery"
    out.mkdir()
    workflow = {"prompt_scheme": "undergraduate_humanities", "thesis_brief": {"title": "Word fixture thesis", "requirements": "Fixture A4"}}
    write_context(root, "ug", workflow)
    tex = submit / "thesis.tex"
    tex.write_text(r"\section{Body}\begin{table}\end{table}\includegraphics{figure.png}", encoding="utf-8")
    (submit / "submission_manifest.json").write_text(json.dumps({"paper_id": "ug", "submit_tex": tex.relative_to(root).as_posix()}))
    figure = submit / "figure.png"
    Image.new("RGB", (400, 400), "white").save(figure)
    docx = out / "thesis.docx"
    doc = Document()
    doc.add_paragraph("Word fixture thesis", "Title")
    doc.add_heading("Body", 1)
    doc.add_paragraph("This is an isolated mechanical test of the delivery validator. " * 5)
    doc.add_picture(str(figure), width=Inches(1))
    doc.add_table(rows=2, cols=2)
    doc.save(docx)
    pdf = out / "preview.pdf"
    with pdfium.PdfDocument.new() as document:
        document.new_page(595, 842)
        document.save(pdf)
    png = out / "page-1.png"
    Image.new("RGB", (600, 840), "white").save(png)
    styles = {key: {"font": "宋体", "size_pt": 12} for key in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "Caption")}
    layout = {"schema": SCHEMA, "unresolved": [], "page": {"width_cm":21,"height_cm":29.7,"top_cm":2.5,"bottom_cm":2.5,"left_cm":3,"right_cm":2.5},
              "styles": styles, "requirements": [{"requirement":"A4", "source":"Fixture", "implementation":"Page"}]}
    layout_path = paper / "word_layout.json"
    layout_path.write_text(json.dumps(layout), encoding="utf-8")
    rec = lambda path: {"path": path.relative_to(root).as_posix(), "sha256": digest(path)}
    receipt = {"schema": SCHEMA, "paper_id":"ug", "status":"awaiting_visual_review", "brief_sha256":brief_digest(workflow),
               "docx":rec(docx), "pdf":rec(pdf), "pages":[rec(png)], "page_count":1,
               "inputs":[rec(tex),rec(figure),rec(layout_path)], "expected":{"headings":1,"tables":1,"figures":1},
               "requirements":layout["requirements"]}
    receipt_path = out / "word_delivery.json"
    receipt_path.write_text(json.dumps(receipt))
    review = {"delivery_sha256":digest(receipt_path), "status":"passed", "issues":[],
              "pages":[{"number":1,"sha256":digest(png),"status":"passed","checked":"Synthetic unit fixture only"}],
              "requirements":[{"number":1,"status":"passed","evidence":"Fixture only"}]}
    (out / "word_review.json").write_text(json.dumps(review))
    return root, paper, workflow


def test_complete_source_bound_word_contract(delivery):
    root, paper, workflow = delivery
    assert validate_word_delivery(root, "ug", workflow)[0] == []
    assert paper_workflow(root, "ug") == workflow


@pytest.mark.parametrize("target", ["delivery/thesis.docx", "delivery/preview.pdf", "delivery/page-1.png", "submit/thesis.tex", "submit/figure.png", "word_layout.json"])
def test_modified_artifact_invalidates_review(delivery, target):
    root, paper, workflow = delivery
    with (paper / target).open("ab") as stream:
        stream.write(b"changed")
    assert validate_word_delivery(root, "ug", workflow)[0]


@pytest.mark.parametrize("target", ["delivery/thesis.docx", "delivery/preview.pdf", "delivery/page-1.png", "delivery/word_review.json"])
def test_missing_required_delivery_blocks_completion(delivery, target):
    root, paper, workflow = delivery
    (paper / target).unlink()
    assert validate_word_delivery(root, "ug", workflow)[0]


def test_changed_brief_invalidates_old_word(delivery):
    root, _, workflow = delivery
    workflow["thesis_brief"]["requirements"] += " changed"
    assert validate_word_delivery(root, "ug", workflow)[0]


@pytest.mark.parametrize("mutation", ["omit_page", "wrong_page_hash", "omit_requirement", "open_issue"])
def test_partial_or_false_review_is_rejected(delivery, mutation):
    root, paper, workflow = delivery
    path = paper / "delivery/word_review.json"
    review = json.loads(path.read_text())
    if mutation == "omit_page": review["pages"] = []
    if mutation == "wrong_page_hash": review["pages"][0]["sha256"] = "wrong"
    if mutation == "omit_requirement": review["requirements"] = []
    if mutation == "open_issue": review["issues"] = ["table clipped"]
    path.write_text(json.dumps(review))
    assert validate_word_delivery(root, "ug", workflow)[0]


def test_fake_docx_cannot_pass_by_changing_receipt_hash(delivery):
    root, paper, workflow = delivery
    path = paper / "delivery/thesis.docx"
    path.write_text("fake renamed TeX")
    receipt_path = paper / "delivery/word_delivery.json"
    receipt = json.loads(receipt_path.read_text())
    receipt["docx"]["sha256"] = digest(path)
    receipt_path.write_text(json.dumps(receipt))
    assert validate_word_delivery(root, "ug", workflow)[0]


def test_tex_only_undergraduate_is_not_finished_but_journal_unchanged(tmp_path):
    from test_figure_sdk_v2 import _write_step3, PAPER_ID
    from unified_paper_orchestrator.figure_output_contract import validate_figure_step
    _write_step3(tmp_path, PAPER_ID)
    assert validate_figure_step(tmp_path, PAPER_ID, 3).ok
    workflow = {"prompt_scheme":"undergraduate_software", "thesis_brief":{"title":"Fixture", "requirements":"A4"}}
    result = validate_figure_step(tmp_path, PAPER_ID, 3, workflow=workflow)
    assert not result.ok and "Word" in " ".join(result.issues)


def test_tex_placeholders_never_enter_formal_word(tmp_path):
    source = tmp_path / "draft.tex"
    source.write_text("测试【验】", encoding="utf-8")
    with pytest.raises(ValueError, match="未核验"):
        _convert_tex(source, tmp_path, tmp_path)


def test_unknown_school_layout_is_not_silently_accepted(delivery):
    _, paper, _ = delivery
    layout = json.loads((paper / "word_layout.json").read_text())
    layout["unresolved"] = ["school section pagination unclear"]
    with pytest.raises(ValueError, match="未解决"):
        _layout(layout)


def test_old_tex_complete_state_does_not_display_word_complete():
    from unified_paper_orchestrator.models import undergraduate_stage_boundary
    from unified_paper_orchestrator.ui_next.read_model import paper_rows, selected_paper
    stages = {"writing":{"status":"complete","prompt_scheme":"undergraduate_software"},"figures":{"status":"complete"}}
    assert undergraduate_stage_boundary(stages, "done") == "figures"
    state = {"papers":{"ug":{"current_stage":"done","stages":stages}}}
    assert not paper_rows(state)[0]["isComplete"]
    assert paper_rows(state)[0]["progress"] < 1
    assert selected_paper(state,"ug")["status"] == "needs_review"
