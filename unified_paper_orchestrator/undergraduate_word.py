"""Local editable Word delivery, with source-bound rendering and review receipts.

No model call is made here. The figure agent supplies the school layout, runs
this exporter, reviews the actual page images, and records its findings.
"""
from __future__ import annotations

import copy
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path
from typing import Any, Mapping
from zipfile import ZipFile, BadZipFile
from xml.etree import ElementTree as ET

from .io_utils import read_json, write_json_atomic

SCHEMA = "undergraduate-word-delivery/v1"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def safe_path(root: Path, value: Any) -> Path:
    path = Path(str(value or ""))
    path = (path if path.is_absolute() else root / path).resolve()
    path.relative_to(root.resolve())
    if not path.is_file():
        raise ValueError(f"交付文件不存在：{path.name}")
    return path


def paper_directory(root: Path, paper_id: str) -> Path:
    if not paper_id or paper_id in {".", ".."} or any(c in paper_id for c in '/\\:'):
        raise ValueError("论文 ID 非法")
    path = (root / "papers" / paper_id).resolve()
    path.relative_to(root.resolve())
    return path


def paper_workflow(root: Path, paper_id: str) -> dict:
    from .config import LEGACY_PROJECTS
    paths = [root / "controller" / "figure_pipeline_state.json"]
    configured = LEGACY_PROJECTS["figures"]
    if Path(configured["root"]).resolve() == root.resolve():
        paths.insert(0, Path(configured["state_path"]))
    for path in paths:
        data = read_json(path, {})
        workflow = data.get("papers", {}).get(paper_id, {}).get("workflow_snapshot", {})
        if workflow:
            return workflow
    context = read_json(paper_directory(root, paper_id) / "undergraduate_delivery.json", {})
    return context.get("workflow_snapshot", {})


def is_word_workflow(workflow: Any) -> bool:
    return isinstance(workflow, Mapping) and workflow.get("prompt_scheme") in {
        "undergraduate_humanities", "undergraduate_software"}


def write_context(root: Path, paper_id: str, workflow: dict) -> None:
    if is_word_workflow(workflow):
        write_json_atomic(paper_directory(root, paper_id) / "undergraduate_delivery.json",
                          {"schema": SCHEMA, "workflow_snapshot": workflow})


def brief_digest(workflow: Mapping) -> str:
    return hashlib.sha256(json.dumps(workflow.get("thesis_brief", {}), ensure_ascii=False,
                                     sort_keys=True).encode("utf-8")).hexdigest()


def pandoc_executable() -> Path:
    base = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parents[1]))
    bundled = base / "runtime" / "pandoc-3.11" / "pandoc-3.11" / "pandoc.exe"
    if bundled.is_file():
        return bundled
    found = shutil.which("pandoc")
    if found:
        return Path(found)
    raise ValueError("Word 转换器缺失，请使用包含 Pandoc 的完整新版软件")


def _run(args: list[str], *, cwd: Path, timeout: int = 180) -> str:
    result = subprocess.run(args, cwd=cwd, capture_output=True, encoding="utf-8", errors="replace",
                            timeout=timeout, creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    if result.returncode:
        raise ValueError("Word 处理失败：" + (result.stderr or result.stdout)[-1800:])
    return result.stdout


def render_word(docx: Path, output: Path) -> tuple[Path, list[Path], str]:
    """Update Word fields and render from the exact DOCX, not the TeX preview."""
    import pypdfium2 as pdfium
    output.mkdir(parents=True, exist_ok=True)
    pdf = output / (docx.stem + ".pdf")
    if os.name == "nt":
        request = output / "word_render_request.json"
        write_json_atomic(request, {"docx": str(docx.resolve()), "pdf": str(pdf.resolve())})
        script = Path(__file__).with_name("word_render.ps1")
        powershell = Path(os.environ.get("SYSTEMROOT", r"C:\Windows")) / "System32/WindowsPowerShell/v1.0/powershell.exe"
        _run([str(powershell), "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass",
              "-File", str(script), "-Request", str(request)], cwd=output)
        renderer = "Microsoft Word"
    else:
        soffice = shutil.which("soffice")
        if not soffice:
            raise ValueError("缺少 Word 排版引擎，不能验收 Word 交付")
        profile = output / "lo-profile"
        _run([soffice, f"-env:UserInstallation={profile.as_uri()}", "--headless", "--convert-to", "pdf",
              "--outdir", str(output), str(docx)], cwd=output)
        renderer = "LibreOffice"
    if not pdf.is_file() or pdf.stat().st_size < 100:
        raise ValueError("Word 没有成功导出 PDF 预览")
    pages = []
    with pdfium.PdfDocument(pdf) as document:
        if not 1 <= len(document) <= 500:
            raise ValueError("Word 预览页数异常")
        for index in range(len(document)):
            page = document[index]
            image = page.render(scale=1.5).to_pil()
            path = output / f"page-{index + 1}.png"
            image.save(path)
            image.close()
            page.close()
            pages.append(path)
    return pdf, pages, renderer


def _layout(value: Any) -> dict:
    if not isinstance(value, dict) or value.get("schema") != SCHEMA:
        raise ValueError("需要 word_layout.json，schema 为 undergraduate-word-delivery/v1")
    result = copy.deepcopy(value)
    if result.get("unresolved") != []:
        raise ValueError("学校格式仍有未解决项，请先处理 word_layout.json.unresolved")
    if not isinstance(result.get("requirements"), list) or not result["requirements"]:
        raise ValueError("缺少逐项学校要求及其来源")
    for item in result["requirements"]:
        if not all(str(item.get(k, "")).strip() for k in ("requirement", "source", "implementation")):
            raise ValueError("学校要求需要 requirement、source、implementation")
    page = result.get("page", {})
    for key in ("width_cm", "height_cm", "top_cm", "bottom_cm", "left_cm", "right_cm"):
        if isinstance(page.get(key), bool) or not isinstance(page.get(key), (int, float)) or not 0 < page[key] < 60:
            raise ValueError(f"学校页面设置无效：{key}")
    if page["left_cm"] + page["right_cm"] >= page["width_cm"] - 4 or page["top_cm"] + page["bottom_cm"] >= page["height_cm"] - 4:
        raise ValueError("学校页边距没有留下足够正文空间")
    styles = result.get("styles", {})
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "Caption"):
        style = styles.get(name, {})
        if not style.get("font") or not 6 <= float(style.get("size_pt", 0)) <= 48:
            raise ValueError(f"缺少明确的 Word 样式：{name}")
    return result


def _convert_tex(source: Path, paper: Path, work: Path) -> tuple[Path, dict, list[Path]]:
    """Resolve local images and common BibTeX/LaTeX references before conversion."""
    import pypdfium2 as pdfium
    text = source.read_text(encoding="utf-8-sig")
    if re.search(r"【验】|SIMULATED_TO_REPLACE|VERIFY_DATA_START", text):
        raise ValueError("整合稿仍有未核验的数据占位，不能生成正式 Word")
    # Inlined input is required: silently dropping an external chapter is unsafe.
    if re.search(r"\\(?:input|include|bibliography)\s*\{", text):
        raise ValueError("请先把章节和参考文献整合进同一 TeX，再生成 Word")
    dependencies = [source]
    def graphic(match):
        raw = match.group(2)
        candidates = [source.parent / raw, paper / raw]
        if not Path(raw).suffix:
            candidates += [p.with_suffix(ext) for p in list(candidates) for ext in (".pdf", ".png", ".jpg")]
        image = next((p.resolve() for p in candidates if p.is_file()), None)
        if image is None:
            raise ValueError(f"Word 缺少图文件：{raw}")
        image.relative_to(paper.resolve())
        dependencies.append(image)
        if image.suffix.lower() == ".pdf":
            target = work / f"image-{digest(image)[:16]}.png"
            with pdfium.PdfDocument(image) as doc:
                if len(doc) != 1:
                    raise ValueError(f"配图 PDF 必须为单页：{image.name}")
                page = doc[0]
                bitmap = page.render(scale=3).to_pil()
                bitmap.save(target)
                bitmap.close()
                page.close()
            image = target
        return r"\includegraphics" + (match.group(1) or "") + "{" + image.as_posix() + "}"
    text = re.sub(r"\\includegraphics(\[[^\]]*\])?\{([^}]+)\}", graphic, text)
    bibkeys = re.findall(r"\\bibitem(?:\[[^\]]*\])?\{([^}]+)\}", text)
    def citation(match):
        keys = [k.strip() for k in match.group(1).split(",")]
        if any(k not in bibkeys for k in keys):
            raise ValueError("Word 转换发现未解析的参考文献引用")
        return "[" + ",".join(str(bibkeys.index(k) + 1) for k in keys) + "]"
    text = re.sub(r"\\cite(?:p|t)?(?:\[[^\]]*\])?\{([^}]+)\}", citation, text)
    text = re.sub(r"\\begin\{thebibliography\}\{[^}]*\}", r"\\section*{参考文献}\n\\begin{enumerate}", text)
    text = re.sub(r"\\bibitem(?:\[[^\]]*\])?\{[^}]+\}", r"\\item ", text)
    text = text.replace(r"\end{thebibliography}", r"\end{enumerate}")
    target = work / "converted.tex"
    target.write_text(text, encoding="utf-8")
    expected = {"figures": len(re.findall(r"\\includegraphics", text)),
                "tables": len(re.findall(r"\\begin\{(?:table\*?|longtable)\}", text)),
                "headings": len(re.findall(r"\\(?:chapter|section|subsection|subsubsection)\*?\{", text))}
    return target, expected, dependencies


def _format_docx(path: Path, layout: dict, title: str) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    doc = Document(path)
    alignments = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
                  "right": WD_ALIGN_PARAGRAPH.RIGHT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}
    for name, options in layout["styles"].items():
        if name not in doc.styles:
            continue
        style = doc.styles[name]
        style.font.name = options.get("latin_font", "Times New Roman")
        style.font.size = Pt(options["size_pt"])
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = bool(options.get("bold", False))
        style.font.italic = False
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        fonts.set(qn("w:eastAsia"), options["font"])
        pf = style.paragraph_format
        pf.space_before, pf.space_after = Pt(options.get("before_pt", 0)), Pt(options.get("after_pt", 0))
        pf.first_line_indent = Pt(options.get("first_line_chars", 0) * options["size_pt"])
        if options.get("line_pt"):
            pf.line_spacing = Pt(options["line_pt"])
        else:
            pf.line_spacing = float(options.get("line_multiple", 1.5))
        pf.alignment = alignments[options.get("align", "left")]
        pf.keep_with_next = name.startswith("Heading") or name == "Caption"
        pf.widow_control = True
    # Pandoc body aliases must not override the requested school body style.
    for name in ("Body Text", "First Paragraph", "Compact"):
        if name in doc.styles:
            doc.styles[name].base_style = doc.styles["Normal"]
            for child in list(doc.styles[name].element):
                if child.tag in {qn("w:pPr"), qn("w:rPr")}:
                    doc.styles[name].element.remove(child)
    for alias, parent in (("Image Caption", "Caption"), ("Table Caption", "Caption"), ("TOC Heading", "Heading 1")):
        if alias in doc.styles:
            doc.styles[alias].base_style = doc.styles[parent]
            for child in list(doc.styles[alias].element):
                if child.tag in {qn("w:pPr"), qn("w:rPr")}:
                    doc.styles[alias].element.remove(child)
            if alias == "TOC Heading":
                level = OxmlElement("w:outlineLvl")
                level.set(qn("w:val"), "9")
                doc.styles[alias].element.get_or_add_pPr().append(level)
    if "Hyperlink" in doc.styles:
        doc.styles["Hyperlink"].font.color.rgb = RGBColor(0, 0, 0)
        doc.styles["Hyperlink"].font.underline = False
    for para in doc.paragraphs:
        kind = {"Image Caption": ("图", "Figure"), "Table Caption": ("表", "Table")}.get(para.style.name)
        if kind:
            prefix = OxmlElement("w:r")
            value = OxmlElement("w:t")
            value.text = kind[0]
            prefix.append(value)
            seq = OxmlElement("w:fldSimple")
            seq.set(qn("w:instr"), "SEQ " + kind[1] + r" \* ARABIC")
            space = OxmlElement("w:r")
            t = OxmlElement("w:t")
            t.set(qn("xml:space"), "preserve")
            t.text = " "
            space.append(t)
            insert = 1 if para._p.pPr is not None else 0
            for offset, node in enumerate((prefix, seq, space)):
                para._p.insert(insert + offset, node)
            para.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.first_line_indent = Pt(0)
            for run in para.runs:
                run.font.italic = False
        if para.text.strip() in layout.get("page_break_before", []):
            para.paragraph_format.page_break_before = True
    p = layout["page"]
    for section in doc.sections:
        for attr, key in (("page_width", "width_cm"), ("page_height", "height_cm"),
                          ("top_margin", "top_cm"), ("bottom_margin", "bottom_cm"),
                          ("left_margin", "left_cm"), ("right_margin", "right_cm")):
            setattr(section, attr, Cm(p[key]))
        section.header_distance = Cm(p.get("header_cm", 1.5))
        section.footer_distance = Cm(p.get("footer_cm", 1.5))
        section.different_first_page_header_footer = bool(layout.get("cover"))
        if "header" in layout:
            section.header.paragraphs[0].text = layout["header"]
        if layout.get("page_numbers", True):
            para = section.footer.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            field = OxmlElement("w:fldSimple")
            field.set(qn("w:instr"), "PAGE")
            para._p.append(field)
    width = Cm(p["width_cm"] - p["left_cm"] - p["right_cm"])
    for shape in doc.inline_shapes:
        if shape.width > width:
            shape.height = int(shape.height * width / shape.width)
            shape.width = width
    for table in doc.tables:
        table.autofit = True
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    para.paragraph_format.first_line_indent = Pt(0)
        header = OxmlElement("w:tblHeader")
        table.rows[0]._tr.get_or_add_trPr().append(header)
    if layout.get("cover"):
        # Generated front matter only contains supplied text; never invented identities.
        first = doc.paragraphs[0] if doc.paragraphs else doc.add_paragraph()
        for index, line in enumerate(layout["cover"]):
            para = first.insert_paragraph_before(str(line), "Title" if index == 0 else "Normal")
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.space_after = Pt(24)
        first.insert_paragraph_before().add_run().add_break(7)
    doc.core_properties.title = title
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.save(path)


def inspect_docx(path: Path, expected: dict, title: str) -> dict:
    with ZipFile(path) as archive:
        xml = ET.fromstring(archive.read("word/document.xml"))
        # Only w:t is visible body text, not field instructions.
        body = "".join(node.text or "" for node in xml.iter(W + "t"))
        figures = len(xml.findall(".//" + W + "drawing"))
        tables = len(xml.findall(".//" + W + "tbl"))
        styles = ET.fromstring(archive.read("word/styles.xml"))
        heading_ids = {node.get(W + "styleId") for node in styles.findall(W + "style")
                       if (level := node.find(W + "pPr/" + W + "outlineLvl")) is not None
                       and level.get(W + "val") in {"0", "1", "2"}}
        headings = len([node for node in xml.iter(W + "pStyle") if node.get(W + "val") in heading_ids])
        if len(body.strip()) < 80 or title not in body:
            raise ValueError("Word 正文缺失或题目不一致")
        if figures < expected.get("figures", 0) or tables < expected.get("tables", 0):
            raise ValueError("Word 丢失图或可编辑表格")
        if headings < expected.get("headings", 0):
            raise ValueError("Word 丢失原生标题结构")
        if re.search(r"【验】|SIMULATED_TO_REPLACE|\\(?:begin|end|includegraphics|cite)\{", body):
            raise ValueError("Word 正文残留数据占位或 LaTeX 源码")
        for name in archive.namelist():
            if name.endswith(".rels"):
                rels = ET.fromstring(archive.read(name))
                if any(r.get("TargetMode") == "External" and r.get("Type", "").endswith(("/image", "/attachedTemplate")) for r in rels):
                    raise ValueError("Word 存在外链图片或模板，不是独立交付文件")
        return {"characters": len(body), "figures": figures, "tables": tables, "headings": headings}


def export_word(root: Path | str, paper_id: str, *, workflow: dict | None = None, render_existing: bool = False) -> dict:
    root = Path(root).resolve()
    paper = paper_directory(root, paper_id)
    workflow = workflow or paper_workflow(root, paper_id)
    if not is_word_workflow(workflow):
        raise ValueError("仅本科论文可使用此 Word 交付流程")
    layout_path = safe_path(paper, "word_layout.json")
    layout = _layout(read_json(layout_path, {}))
    manifest = read_json(paper / "submit" / "submission_manifest.json", {})
    source = safe_path(root, manifest.get("submit_tex"))
    source.relative_to(paper)
    output = paper / "delivery"
    output.mkdir(exist_ok=True)
    receipt = output / "word_delivery.json"
    # Invalidate the previous review before changing any delivery artifact.
    write_json_atomic(receipt, {"schema": SCHEMA, "status": "building", "paper_id": paper_id})
    with tempfile.TemporaryDirectory(prefix="word-build-", dir=paper) as tmp:
        work = Path(tmp)
        tex, expected, dependencies = _convert_tex(source, paper, work)
        docx = work / (paper_id + ".docx")
        args = [str(pandoc_executable()), str(tex), "--from=latex", "--to=docx", "--standalone",
                "--resource-path=" + str(source.parent), "--output=" + str(docx)]
        if layout.get("toc", True):
            args += ["--toc", "--toc-depth=3", "--metadata=toc-title:目录"]
        if layout.get("number_sections", True):
            args += ["--number-sections"]
        if layout.get("reference_docx"):
            reference = safe_path(paper, layout["reference_docx"])
            dependencies.append(reference)
            args += ["--reference-doc=" + str(reference)]
        destination = output / docx.name
        if render_existing:
            safe_path(output, docx.name)
        else:
            _run(args, cwd=work)
            _format_docx(docx, layout, workflow["thesis_brief"]["title"])
            if destination.exists():
                archive = output / "history" / digest(destination)[:16]
                archive.mkdir(parents=True, exist_ok=True)
                shutil.copy2(destination, archive / destination.name)
            shutil.copy2(docx, destination)
    try:
        pdf, pages, renderer = render_word(destination, output / "preview")
        counts = inspect_docx(destination, expected, workflow["thesis_brief"]["title"])
        def record(path):
            return {"path": path.relative_to(root).as_posix(), "sha256": digest(path)}
        result = {"schema": SCHEMA, "paper_id": paper_id, "status": "awaiting_visual_review",
                  "brief_sha256": brief_digest(workflow), "docx": record(destination), "pdf": record(pdf),
                  "pages": [record(path) for path in pages], "page_count": len(pages), "renderer": renderer,
                  "inputs": [record(path) for path in dict.fromkeys([*dependencies, layout_path])],
                  "expected": expected, "counts": counts, "requirements": layout["requirements"]}
        write_json_atomic(receipt, result)
        return result
    except Exception as exc:
        write_json_atomic(receipt, {"schema": SCHEMA, "paper_id": paper_id, "status": "failed", "error": str(exc)})
        raise


def validate_word_delivery(root: Path, paper_id: str, workflow: dict) -> tuple[list[str], list[Path]]:
    paper = paper_directory(root, paper_id)
    receipt_path = paper / "delivery" / "word_delivery.json"
    review_path = paper / "delivery" / "word_review.json"
    artifacts = [p for p in (receipt_path, review_path) if p.is_file()]
    try:
        receipt = read_json(receipt_path, {})
        if receipt.get("schema") != SCHEMA or receipt.get("status") != "awaiting_visual_review" or receipt.get("paper_id") != paper_id:
            raise ValueError("本科论文尚未生成并渲染 Word 正式稿")
        if receipt.get("brief_sha256") != brief_digest(workflow):
            raise ValueError("Word 使用的题目/学校材料版本已变化")
        layout_path = safe_path(paper, "word_layout.json")
        layout = _layout(read_json(layout_path, {}))
        manifest = read_json(paper / "submit" / "submission_manifest.json", {})
        source = safe_path(root, manifest.get("submit_tex"))
        recorded_inputs = {safe_path(root, item["path"]) for item in receipt.get("inputs", [])}
        if source not in recorded_inputs or layout_path not in recorded_inputs or receipt.get("requirements") != layout["requirements"]:
            raise ValueError("Word 验收未绑定当前整合稿与学校格式配置")
        source_text = source.read_text(encoding="utf-8-sig")
        expected = {"figures": len(re.findall(r"\\includegraphics", source_text)),
                    "tables": len(re.findall(r"\\begin\{(?:table\*?|longtable)\}", source_text)),
                    "headings": len(re.findall(r"\\(?:chapter|section|subsection|subsubsection)\*?\{", source_text))}
        pages = receipt.get("pages", [])
        if not pages or receipt.get("page_count") != len(pages):
            raise ValueError("Word 缺少完整逐页预览")
        for item in [receipt["docx"], receipt["pdf"], *pages, *receipt["inputs"]]:
            path = safe_path(root, item["path"])
            path.relative_to(paper)
            if digest(path) != item["sha256"]:
                raise ValueError("Word、配图、输入或预览已变化，需要重新生成并复核")
            artifacts.append(path)
        import pypdfium2 as pdfium
        from PIL import Image
        with pdfium.PdfDocument(safe_path(root, receipt["pdf"]["path"])) as document:
            if len(document) != len(pages):
                raise ValueError("Word PDF 与逐页预览数量不一致")
        for item in pages:
            with Image.open(safe_path(root, item["path"])) as image:
                if image.width < 300 or image.height < 300:
                    raise ValueError("Word 页面预览尺寸不足以检查排版")
                image.verify()
        docx = safe_path(root, receipt["docx"]["path"])
        if docx.suffix.lower() != ".docx":
            raise ValueError("正式交付必须是 DOCX")
        inspect_docx(docx, expected, workflow["thesis_brief"]["title"])
        review = read_json(review_path, {})
        if review.get("delivery_sha256") != digest(receipt_path):
            raise ValueError("缺少当前 Word 版本的逐页排版复核")
        if review.get("status") != "passed" or review.get("issues") != []:
            raise ValueError("Word 排版复核尚未通过")
        reviewed = review.get("pages", [])
        if len(reviewed) != len(pages) or any(item.get("number") != index or item.get("sha256") != pages[index-1]["sha256"]
                or item.get("status") != "passed" or not str(item.get("checked", "")).strip()
                for index, item in enumerate(reviewed, 1)):
            raise ValueError("Word 排版复核必须覆盖当前全部页面")
        requirements = review.get("requirements", [])
        if len(requirements) != len(receipt["requirements"]) or any(item.get("number") != index or
                item.get("status") != "passed" or not str(item.get("evidence", "")).strip()
                for index, item in enumerate(requirements, 1)):
            raise ValueError("学校格式要求尚未逐项复核")
        return [], list(dict.fromkeys(artifacts))
    except (OSError, ValueError, KeyError, TypeError, AttributeError, BadZipFile, ET.ParseError, RuntimeError) as exc:
        return [str(exc)], artifacts


def word_prompt(root: Path, paper_id: str) -> str:
    if getattr(sys, "frozen", False):
        command = [sys.executable, "undergraduate-word-export"]
    else:
        command = [sys.executable, "-m", "unified_paper_orchestrator", "undergraduate-word-export"]
    command += ["--root", str(root), "--paper-id", paper_id]
    # PowerShell literal quoting, never JSON-as-shell-escaping.
    quoted = "& " + " ".join("'" + arg.replace("'", "''") + "'" for arg in command)
    return "\n## Word 正式交付（本科画图最终步必须完成）\n" + (
        "已有旧快照中关于只交付 TeX/PDF 的表述已被本次 Word 交付要求替代，写作与数据核验仍沿用本论文快照。"
        "TeX 是内部中间稿。图表整合后必须生成可编辑 Word、由 Word 导出的 PDF 预览和逐页复核。"
        "先按学校资料编写 papers/<ID>/word_layout.json，禁止凭空声称套用模板。"
        "格式结构见下方示例；示例的具体值不是学校标准，必须按实际要求调整，未规定项注明采用通用值。"
        "复杂封面/分节可使用本地 Word 工具进一步排版。可直接编辑 delivery/<ID>.docx；修改后用下方命令追加 "
        "--render-existing 重新渲染现有 Word、更新交付记录，不会重新转换并覆盖修改。不得伪造哈希或验收。\n"
        '{"schema":"undergraduate-word-delivery/v1","unresolved":[],"requirements":[{"requirement":"具体要求",'
        '"source":"学校资料文件及页码或未规定说明","implementation":"对应格式配置及位置"}],'
        '"page":{"width_cm":21,"height_cm":29.7,"top_cm":2.5,"bottom_cm":2.5,"left_cm":3,"right_cm":2.5},'
        '"styles":{"Normal":{"font":"宋体","size_pt":12,"first_line_chars":2,"line_multiple":1.5},'
        '"Title":{"font":"黑体","size_pt":20,"align":"center"},"Heading 1":{"font":"黑体","size_pt":16},'
        '"Heading 2":{"font":"黑体","size_pt":14},"Heading 3":{"font":"黑体","size_pt":12},'
        '"Caption":{"font":"宋体","size_pt":10.5,"align":"center"}},"toc":true,"page_numbers":true}\n'
        "可选 reference_docx 为本论文目录内的学校样式参考 DOCX；它只继承样式，不自动复制学校封面。"
        "可选 cover 为已有真实信息组成的封面行，header 为页眉。"
        "可选 page_break_before 为需要另起一页的标题文字数组，number_sections 控制章节编号。"
        "不支持的特殊学校要求必须解决或报告阻塞，不能填 unresolved=[] 冒充完成。\n"
        "调用以下本地转换命令（不会启动模型或投稿）：\n" + quoted + "\n"
        "读取 delivery/word_delivery.json，逐张打开其 pages 中的 PNG，核对中文、目录/页码、标题、图表、"
        "引用、分页、遮挡与全部学校要求。发现问题修改后重新导出并重看；全部通过才写 delivery/word_review.json："
        '{"delivery_sha256":"word_delivery.json当前SHA256","status":"passed","issues":[],"pages":'
        '[{"number":1,"sha256":"对应PNG的SHA256","status":"passed","checked":"本页实际检查内容"}],'
        '"requirements":[{"number":1,"status":"passed","evidence":"实际验证位置与结果"}]}。'
        "pages 和 requirements 必须逐项完整记录；禁止未看页面就宣称通过。最终向用户交付 delivery/ 中的 DOCX。\n"
    ).replace("<ID>", paper_id)
